"""Run with:  python -m unittest discover -s finance_tracker/tests"""
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import ingest  # noqa: E402

SAMPLES = ROOT / "samples"


def run(*argv):
    assert ingest.main([str(a) for a in argv]) == 0


def txn_rows(tracker):
    wb = load_workbook(tracker)
    return ingest.read_transactions(wb)


class ParseTests(unittest.TestCase):
    def test_amounts(self):
        cases = {"1,234.56": 1234.56, "-$12.00": -12.0, "(12.00)": -12.0, "$5": 5.0,
                 "12.00 CR": 12.0, "12.00 DR": -12.0, "12.00-": -12.0, "": None, "abc": None}
        for text, want in cases.items():
            self.assertEqual(ingest.parse_amount(text), want, text)

    def test_csv_with_preamble_and_debit_credit_columns(self):
        txns = ingest.parse_csv(SAMPLES / "sample_checking.csv")
        self.assertEqual(len(txns), 11)  # "Ending Balance" footer is skipped
        self.assertEqual(txns[0], ingest.Txn(date(2026, 9, 1), "PAYROLL ACME CORP DIRECT DEP", 4250.0))
        self.assertEqual(txns[1].amount, -2100.0)

    def test_csv_with_single_amount_column(self):
        txns = ingest.parse_csv(SAMPLES / "sample_card.csv")
        self.assertEqual([t.amount for t in txns], [-15.49, -23.80, -64.12, 19.99, 1312.40])
        self.assertEqual(txns[0].date, date(2026, 10, 1))  # Transaction Date, not Post Date

    def test_day_first_dates(self):
        self.assertEqual(ingest.parse_date("03/04/2026", day_first=True), date(2026, 4, 3))
        self.assertEqual(ingest.parse_date("03/04/2026"), date(2026, 3, 4))

    def test_pdf_infers_year_across_new_year(self):
        txns = ingest.parse_pdf(SAMPLES / "sample_card_statement.pdf")
        self.assertEqual(len(txns), 6)
        self.assertEqual(txns[0].date, date(2026, 12, 8))
        self.assertEqual(txns[-1].date, date(2027, 1, 3))
        self.assertEqual(txns[3].amount, -500.0)

    def test_categorize_matches_word_starts(self):
        rules = [("fee", "Fees"), ("uber eats", "Dining"), ("uber", "Transport")]
        self.assertEqual(ingest.categorize("MONTHLY FEE", rules), "Fees")
        self.assertEqual(ingest.categorize("BLUE BOTTLE COFFEE", rules), ingest.UNCATEGORIZED)
        self.assertEqual(ingest.categorize("UBER EATS 123", rules), "Dining")
        self.assertEqual(ingest.categorize("UBER *TRIP", rules), "Transport")


class TrackerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.tracker = Path(self.tmp.name) / "tracker.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_reimport_skips_duplicates_but_keeps_same_day_repeats(self):
        csv = SAMPLES / "sample_checking.csv"
        run(csv, "--account", "Checking", "--tracker", self.tracker)
        run(csv, "--account", "Checking", "--tracker", self.tracker)
        rows = txn_rows(self.tracker)
        self.assertEqual(len(rows), 11)
        starbucks = [r for r in rows if "STARBUCKS" in r[ingest.COL["Description"]]]
        self.assertEqual(len(starbucks), 2)  # two identical coffees on 09/05 both kept
        imports = list(load_workbook(self.tracker)[ingest.IMPORTS_SHEET].values)
        self.assertEqual(imports[-1][4:], (0, 11))

    def test_flip_sign_and_summary_tabs(self):
        run(SAMPLES / "sample_card_statement.pdf", "--account", "Visa", "--flip-sign",
            "--tracker", self.tracker)
        wb = load_workbook(self.tracker)
        rows = ingest.read_transactions(wb)
        self.assertEqual(rows[0][ingest.COL["Amount"]], -54.21)
        monthly = [r[0] for r in wb[ingest.MONTHLY_SHEET].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(monthly[:2], ["2026-12", "2027-01"])
        self.assertEqual(wb[ingest.CATEGORY_SHEET]["A2"].value, "Card Payment")
        self.assertEqual(wb[ingest.ACCOUNT_SHEET]["A2"].value, "Visa")

    def test_manual_categories_survive_and_recategorize(self):
        run(SAMPLES / "sample_checking.csv", "--account", "Checking", "--tracker", self.tracker)
        wb = load_workbook(self.tracker)
        ws = wb[ingest.TXN_SHEET]
        for row in ws.iter_rows(min_row=2):
            desc = row[ingest.COL["Description"]].value
            if desc.startswith("RENT"):
                row[ingest.COL["Category"]].value = "Home"  # user override
        wb[ingest.RULES_SHEET].append(["local hardware", "Home Improvement"])
        wb.save(self.tracker)

        run("--recategorize", "--tracker", self.tracker)
        cats = {r[ingest.COL["Description"]]: r[ingest.COL["Category"]] for r in txn_rows(self.tracker)}
        self.assertEqual(cats["RENT PAYMENT - OAK APARTMENTS"], "Home")
        self.assertEqual(cats["LOCAL HARDWARE CO"], "Home Improvement")

    def test_dry_run_does_not_write(self):
        run(SAMPLES / "sample_card.csv", "--account", "Visa", "--tracker", self.tracker, "--dry-run")
        self.assertFalse(self.tracker.exists())


if __name__ == "__main__":
    unittest.main()
