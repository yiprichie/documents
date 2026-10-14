#!/usr/bin/env python3
"""Append financial statements (CSV or PDF) to a tracking spreadsheet.

Usage examples:
    python ingest.py --init                                   # create tracker.xlsx
    python ingest.py statements/chase_2026-09.csv --account "Chase Checking"
    python ingest.py amex_sep.pdf --account "Amex" --flip-sign
    python ingest.py statements/*.csv --account "Chase Checking" --dry-run

Each run:
  * parses every statement into (date, description, amount) rows,
  * skips rows already in the tracker (so re-importing a file is safe),
  * categorizes new rows with the keyword rules on the Categories tab,
  * rebuilds the summary tabs (Monthly Summary, By Category, By Account),
  * logs the import on the Imports tab.

Amount convention in the tracker: money in is positive, money out is negative.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
DEFAULT_TRACKER = HERE / "tracker.xlsx"
DEFAULT_RULES = HERE / "categories.csv"
UNCATEGORIZED = "Uncategorized"
# Money moving between your own accounts. Left out of Income/Spending on the
# Monthly Summary so a card payment isn't counted once in checking and again on the card.
INTERNAL_CATEGORIES = ["Card Payment", "Transfers"]

TXN_SHEET = "Transactions"
MONTHLY_SHEET = "Monthly Summary"
CATEGORY_SHEET = "By Category"
ACCOUNT_SHEET = "By Account"
IMPORTS_SHEET = "Imports"
RULES_SHEET = "Categories"
SHEET_ORDER = [MONTHLY_SHEET, CATEGORY_SHEET, ACCOUNT_SHEET, TXN_SHEET, IMPORTS_SHEET, RULES_SHEET]

TXN_COLUMNS = ["Date", "Month", "Account", "Description", "Category", "Amount",
               "Notes", "Source File", "Imported At", "Txn ID"]
COL = {name: i for i, name in enumerate(TXN_COLUMNS)}
IMPORT_COLUMNS = ["Imported At", "Source File", "Account", "Rows Parsed",
                  "Rows Added", "Duplicates Skipped"]
RULE_COLUMNS = ["Keyword", "Category"]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TOTAL_FONT = Font(bold=True)
MONEY = '#,##0.00;[Red]-#,##0.00'


@dataclass
class Txn:
    date: date
    description: str
    amount: float


# --------------------------------------------------------------------------
# Value parsing
# --------------------------------------------------------------------------

DATE_FORMATS_MDY = ["%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%m-%d-%y"]
DATE_FORMATS_DMY = ["%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y"]
DATE_FORMATS_COMMON = ["%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%d %b %Y", "%d %B %Y",
                       "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%d-%b-%Y", "%d-%b-%y"]


def parse_date(text: str, day_first: bool = False, fmt: str | None = None) -> date | None:
    text = text.strip()
    if not text:
        return None
    formats = [fmt] if fmt else (
        DATE_FORMATS_COMMON + (DATE_FORMATS_DMY if day_first else DATE_FORMATS_MDY))
    for f in formats:
        try:
            return datetime.strptime(text, f).date()
        except ValueError:
            continue
    # Some exports append a time ("2026-09-01 00:00:00").
    if " " in text and not fmt:
        return parse_date(text.split(" ")[0], day_first)
    return None


AMOUNT_RE = re.compile(r"^\(?-?\$?-?[\d,]*\.?\d+\)?$")


def parse_amount(text: str) -> float | None:
    """Parse '1,234.56', '-$12.00', '(12.00)', '12.00 CR', '12.00-' into a float."""
    s = str(text).strip().replace("−", "-").replace(" ", "")
    if not s:
        return None
    sign = 1
    upper = s.upper()
    if upper.endswith("CR"):
        s = s[:-2]
    elif upper.endswith("DR"):
        s, sign = s[:-2], -1
    if s.endswith("-"):
        s, sign = s[:-1], -sign
    for cur in ("USD", "CAD", "EUR", "GBP", "£", "€"):
        s = s.replace(cur, "")
    if not AMOUNT_RE.match(s):
        return None
    if s.startswith("(") and s.endswith(")"):
        s, sign = s[1:-1], -sign
    s = s.replace("$", "").replace(",", "")
    if s.count("-") > 1:
        return None
    try:
        return round(sign * float(s), 2)
    except ValueError:
        return None


# --------------------------------------------------------------------------
# CSV statements
# --------------------------------------------------------------------------

HEADER_ALIASES = {
    "date": ["date", "transaction date", "trans date", "trans. date", "posted date",
             "posting date", "post date", "value date", "booking date"],
    "description": ["description", "transaction description", "payee", "merchant",
                    "details", "name", "memo", "narrative", "reference"],
    "amount": ["amount", "transaction amount", "amount (usd)", "amount usd", "value"],
    "debit": ["debit", "debits", "withdrawal", "withdrawals", "money out", "debit amount",
              "paid out", "outflow"],
    "credit": ["credit", "credits", "deposit", "deposits", "money in", "credit amount",
               "paid in", "inflow"],
}


def _norm(header: str) -> str:
    return re.sub(r"\s+", " ", header.strip().lower().lstrip("﻿"))


def _map_headers(row: list[str]) -> dict[str, int]:
    found: dict[str, int] = {}
    normed = [_norm(h) for h in row]
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:  # alias order is preference order
            if alias in normed and normed.index(alias) not in found.values():
                found[field] = normed.index(alias)
                break
    return found


def parse_csv(path: Path, day_first: bool = False, date_format: str | None = None) -> list[Txn]:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(raw), dialect))

    # Banks often put a few lines of account info above the real header row.
    header_idx, mapping = None, {}
    for i, row in enumerate(rows[:30]):
        m = _map_headers(row)
        if "date" in m and "description" in m and ("amount" in m or "debit" in m or "credit" in m):
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise ValueError(
            f"{path.name}: couldn't find a header row with date, description and amount "
            f"columns. Recognized names: {HEADER_ALIASES}")

    txns = []
    for row in rows[header_idx + 1:]:
        if not any(c.strip() for c in row):
            continue
        cell = lambda f: row[mapping[f]] if f in mapping and mapping[f] < len(row) else ""
        d = parse_date(cell("date"), day_first, date_format)
        if d is None:
            continue  # footer / totals line
        if "amount" in mapping:
            amt = parse_amount(cell("amount"))
        else:
            debit, credit = parse_amount(cell("debit")), parse_amount(cell("credit"))
            if debit is None and credit is None:
                amt = None
            else:
                amt = abs(credit or 0) - abs(debit or 0)
        if amt is None:
            continue
        txns.append(Txn(d, " ".join(cell("description").split()), amt))
    return txns


# --------------------------------------------------------------------------
# PDF statements
# --------------------------------------------------------------------------

MONTHS = "Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
_DATE_PAT = (rf"(?:\d{{4}}-\d{{2}}-\d{{2}}|\d{{1,2}}[/-]\d{{1,2}}(?:[/-]\d{{2,4}})?"
             rf"|(?:{MONTHS})[a-z]*\.? \d{{1,2}}(?:,? \d{{4}})?|\d{{1,2}} (?:{MONTHS})[a-z]*(?: \d{{4}})?)")
_AMT_PAT = r"(?:-?\(?-?\$?\d{1,3}(?:,\d{3})*\.\d{2}\)?-?(?:\s?(?:CR|DR))?)"
PDF_LINE_RE = re.compile(
    rf"^(?P<date>{_DATE_PAT})\s+(?:(?:{_DATE_PAT})\s+)?"   # optional second (post) date
    rf"(?P<desc>.+?)\s+(?P<amount>{_AMT_PAT})(?:\s+(?P<balance>{_AMT_PAT}))?$",
    re.IGNORECASE)
FULL_DATE_RE = re.compile(
    rf"\b(\d{{1,2}}/\d{{1,2}}/\d{{4}}|\d{{4}}-\d{{2}}-\d{{2}}|(?:{MONTHS})[a-z]* \d{{1,2}},? \d{{4}})\b",
    re.IGNORECASE)


def _statement_end_date(text: str, day_first: bool) -> date | None:
    """Best guess at the statement closing date: the latest full date printed on it."""
    dates = []
    for m in FULL_DATE_RE.findall(text):
        d = parse_date(m.replace(",", ", ").replace(",  ", ", "), day_first) or parse_date(m, day_first)
        if d:
            dates.append(d)
    return max(dates) if dates else None


def _parse_pdf_date(text: str, end: date | None, day_first: bool, year: int | None) -> date | None:
    text = re.sub(r"\s+", " ", text.replace(".", "")).strip()
    d = parse_date(text, day_first)
    if d:
        return d
    # Dates without a year ("09/14", "Sep 14", "14 Sep"): borrow the statement's year,
    # rolling back a year for December rows on a January statement.
    ref_year = year or (end.year if end else date.today().year)
    for fmt in (["%d/%m", "%d-%m"] if day_first else ["%m/%d", "%m-%d"]) + ["%b %d", "%B %d", "%d %b", "%d %B"]:
        try:
            d = datetime.strptime(f"{text} {ref_year}", f"{fmt} %Y").date()
        except ValueError:
            continue
        if end and not year and d > end and (d - end).days > 31:
            try:
                d = d.replace(year=ref_year - 1)
            except ValueError:  # Feb 29 in a non-leap year
                return None
        return d
    return None


def parse_pdf(path: Path, day_first: bool = False, year: int | None = None) -> list[Txn]:
    import pdfplumber  # imported lazily so CSV-only users don't need it

    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    if not text.strip():
        raise ValueError(f"{path.name}: no text found. It may be a scanned image; "
                         "OCR it first or download a CSV export instead.")
    end = _statement_end_date(text, day_first)
    txns = []
    for line in text.splitlines():
        line = " ".join(line.split())
        m = PDF_LINE_RE.match(line)
        if not m:
            continue
        d = _parse_pdf_date(m["date"], end, day_first, year)
        amt = parse_amount(m["amount"])
        desc = m["desc"].strip()
        if d is None or amt is None or not desc:
            continue
        txns.append(Txn(d, desc, amt))
    return txns


def parse_statement(path: Path, args) -> list[Txn]:
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        txns = parse_csv(path, args.day_first, args.date_format)
    elif suffix == ".pdf":
        txns = parse_pdf(path, args.day_first, args.year)
    else:
        raise ValueError(f"{path.name}: unsupported file type {suffix!r} (use .csv or .pdf)")
    if args.flip_sign:
        txns = [Txn(t.date, t.description, -t.amount) for t in txns]
    return txns


# --------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------

def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _set_widths(ws, widths: dict[str, int]) -> None:
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w


def load_rules_csv(path: Path) -> list[tuple[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [(r["Keyword"].strip(), r["Category"].strip())
                for r in csv.DictReader(f) if r.get("Keyword", "").strip()]


def new_workbook(rules: list[tuple[str, str]]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEET_ORDER:
        wb.create_sheet(name)

    ws = wb[TXN_SHEET]
    ws.append(TXN_COLUMNS)
    _style_header(ws, len(TXN_COLUMNS))
    _set_widths(ws, {"A": 12, "B": 9, "C": 18, "D": 48, "E": 18, "F": 12,
                     "G": 24, "H": 28, "I": 18, "J": 14})
    ws.auto_filter.ref = "A1:J1"

    ws = wb[IMPORTS_SHEET]
    ws.append(IMPORT_COLUMNS)
    _style_header(ws, len(IMPORT_COLUMNS))
    _set_widths(ws, {"A": 18, "B": 32, "C": 18, "D": 12, "E": 12, "F": 18})

    ws = wb[RULES_SHEET]
    ws.append(RULE_COLUMNS)
    _style_header(ws, 2)
    _set_widths(ws, {"A": 28, "B": 22, "D": 70})
    for kw, cat in rules:
        ws.append([kw, cat])
    ws["D1"] = "How rules work"
    ws["D1"].font = TOTAL_FONT
    ws["D2"] = ("New transactions get the Category of the first row whose Keyword appears "
                "in the description (case-insensitive).")
    ws["D3"] = ("Edit or add rows here freely. A category you type by hand on the "
                "Transactions tab is never overwritten.")
    ws["D4"] = "Run ingest.py with --recategorize to re-apply rules to Uncategorized rows."

    rebuild_summaries(wb)
    return wb


def read_rules(wb) -> list[tuple[str, str]]:
    ws = wb[RULES_SHEET]
    return [(str(kw).strip().lower(), str(cat).strip())
            for kw, cat in ws.iter_rows(min_row=2, max_col=2, values_only=True)
            if kw and cat]


def categorize(description: str, rules: list[tuple[str, str]]) -> str:
    d = description.lower()
    for kw, cat in rules:
        # Keyword must start at a word boundary, so "fee" won't match "coffee".
        if re.search(r"(?<![a-z0-9])" + re.escape(kw), d):
            return cat
    return UNCATEGORIZED


def read_transactions(wb) -> list[list]:
    ws = wb[TXN_SHEET]
    rows = []
    for r in ws.iter_rows(min_row=2, max_col=len(TXN_COLUMNS), values_only=True):
        if r[COL["Txn ID"]]:
            row = list(r)
            if isinstance(row[COL["Date"]], datetime):
                row[COL["Date"]] = row[COL["Date"]].date()
            rows.append(row)
    return rows


def write_transactions(wb, rows: list[list]) -> None:
    ws = wb[TXN_SHEET]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    rows.sort(key=lambda r: (r[COL["Date"]], str(r[COL["Account"]]), str(r[COL["Description"]])))
    for row in rows:
        ws.append(row)
    for r in range(2, len(rows) + 2):
        ws.cell(row=r, column=COL["Date"] + 1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=COL["Amount"] + 1).number_format = MONEY
    last = max(len(rows) + 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TXN_COLUMNS))}{last}"


def txn_id(account: str, t: Txn, occurrence: int) -> str:
    key = f"{account.lower()}|{t.date.isoformat()}|{t.description.lower()}|{t.amount:.2f}|{occurrence}"
    return hashlib.sha1(key.encode()).hexdigest()[:12]


def _reset_sheet(wb, name: str):
    idx = wb.sheetnames.index(name)
    wb.remove(wb[name])
    return wb.create_sheet(name, idx)


def rebuild_summaries(wb) -> None:
    """Rewrite the summary tabs as live formulas over the Transactions tab."""
    rows = read_transactions(wb)
    months = sorted({r[COL["Month"]] for r in rows})
    categories = sorted({r[COL["Category"]] or UNCATEGORIZED for r in rows})
    accounts = sorted({r[COL["Account"]] for r in rows})

    T = f"'{TXN_SHEET}'!"
    month_rng, acct_rng = f"{T}$B:$B", f"{T}$C:$C"
    cat_rng, amt_rng, date_rng = f"{T}$E:$E", f"{T}$F:$F", f"{T}$A:$A"

    # Monthly Summary
    ws = _reset_sheet(wb, MONTHLY_SHEET)
    ws.append(["Month", "Income", "Spending", "Net", "Transactions"])
    _style_header(ws, 5)
    internal = "".join(f',{cat_rng},"<>{c}"' for c in INTERNAL_CATEGORIES)
    for i, m in enumerate(months, start=2):
        ws.append([m,
                   f'=SUMIFS({amt_rng},{month_rng},$A{i},{amt_rng},">0"{internal})',
                   f'=SUMIFS({amt_rng},{month_rng},$A{i},{amt_rng},"<0"{internal})',
                   f"=B{i}+C{i}",
                   f"=COUNTIFS({month_rng},$A{i})"])
    if months:
        n = len(months) + 1
        ws.append(["Total"] + [f"=SUM({c}2:{c}{n})" for c in "BCDE"])
        ws.append([])
        ws.append(["Average / month", f"=AVERAGE(B2:B{n})", f"=AVERAGE(C2:C{n})",
                   f"=AVERAGE(D2:D{n})", f"=AVERAGE(E2:E{n})"])
        for r in (n + 1, n + 3):
            for c in range(1, 6):
                ws.cell(row=r, column=c).font = TOTAL_FONT
    else:
        ws.append(["(no transactions yet; run ingest.py on a statement)"])
    ws.cell(row=ws.max_row + 2, column=1,
            value="Income and Spending leave out these categories (money moved between your "
                  "own accounts): " + ", ".join(INTERNAL_CATEGORIES)).font = Font(italic=True)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=4):
        for c in row:
            c.number_format = MONEY
    _set_widths(ws, {"A": 18, "B": 14, "C": 14, "D": 14, "E": 13})

    # By Category: categories x months
    ws = _reset_sheet(wb, CATEGORY_SHEET)
    ws.append(["Category"] + months + ["Total"])
    _style_header(ws, len(months) + 2)
    for i, cat in enumerate(categories, start=2):
        line = [cat]
        for j, _ in enumerate(months, start=2):
            col = get_column_letter(j)
            line.append(f"=SUMIFS({amt_rng},{cat_rng},$A{i},{month_rng},{col}$1)")
        line.append(f"=SUM(B{i}:{get_column_letter(len(months) + 1)}{i})")
        ws.append(line)
    if categories:
        n = len(categories) + 1
        ws.append(["Total"] + [f"=SUM({get_column_letter(j)}2:{get_column_letter(j)}{n})"
                               for j in range(2, len(months) + 3)])
        for c in range(1, len(months) + 3):
            ws.cell(row=n + 1, column=c).font = TOTAL_FONT
    for row in ws.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = MONEY
    ws.column_dimensions["A"].width = 22
    for j in range(2, len(months) + 3):
        ws.column_dimensions[get_column_letter(j)].width = 12

    # By Account
    ws = _reset_sheet(wb, ACCOUNT_SHEET)
    ws.append(["Account", "First Date", "Last Date", "Money In", "Money Out", "Net", "Transactions"])
    _style_header(ws, 7)
    for i, a in enumerate(accounts, start=2):
        ws.append([a,
                   f"=_xlfn.MINIFS({date_rng},{acct_rng},$A{i})",
                   f"=_xlfn.MAXIFS({date_rng},{acct_rng},$A{i})",
                   f'=SUMIFS({amt_rng},{acct_rng},$A{i},{amt_rng},">0")',
                   f'=SUMIFS({amt_rng},{acct_rng},$A{i},{amt_rng},"<0")',
                   f"=D{i}+E{i}",
                   f"=COUNTIFS({acct_rng},$A{i})"])
        ws.cell(row=i, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=3).number_format = "yyyy-mm-dd"
        for c in (4, 5, 6):
            ws.cell(row=i, column=c).number_format = MONEY
    _set_widths(ws, {"A": 22, "B": 12, "C": 12, "D": 14, "E": 14, "F": 14, "G": 13})


def ingest(paths: list[Path], account: str, tracker: Path, args) -> int:
    if tracker.exists():
        wb = load_workbook(tracker)
    else:
        wb = new_workbook(load_rules_csv(args.rules))
        if not args.dry_run:
            print(f"Creating new tracker at {tracker}")
    rules = read_rules(wb)
    rows = read_transactions(wb)
    seen = {r[COL["Txn ID"]] for r in rows}

    if args.recategorize:
        changed = 0
        for r in rows:
            if (r[COL["Category"]] or UNCATEGORIZED) == UNCATEGORIZED:
                new = categorize(str(r[COL["Description"]]), rules)
                if new != UNCATEGORIZED:
                    r[COL["Category"]], changed = new, changed + 1
        print(f"Recategorized {changed} previously uncategorized rows")

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    log, total_added = [], 0
    for path in paths:
        txns = parse_statement(path, args)
        occurrences: Counter = Counter()
        added = dupes = 0
        for t in txns:
            key = (t.date, t.description.lower(), t.amount)
            occurrences[key] += 1
            tid = txn_id(account, t, occurrences[key])
            if tid in seen:
                dupes += 1
                continue
            seen.add(tid)
            rows.append([t.date, t.date.strftime("%Y-%m"), account, t.description,
                         categorize(t.description, rules), t.amount, None, path.name, now, tid])
            added += 1
        total_added += added
        log.append([now, path.name, account, len(txns), added, dupes])
        print(f"{path.name}: parsed {len(txns)}, added {added}, skipped {dupes} duplicates")
        if args.dry_run:
            for t in txns:
                print(f"    {t.date}  {t.amount:>10.2f}  {t.description}")
        if not txns:
            print(f"    warning: no transactions recognized in {path.name}; "
                  "see README 'When a statement doesn't parse'", file=sys.stderr)

    if args.dry_run:
        print("Dry run: tracker not modified")
        return total_added

    write_transactions(wb, rows)
    for entry in log:
        wb[IMPORTS_SHEET].append(entry)
    rebuild_summaries(wb)
    wb.active = wb.sheetnames.index(MONTHLY_SHEET)
    wb.save(tracker)
    print(f"Saved {tracker} ({len(rows)} transactions total)")
    return total_added


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("statements", nargs="*", type=Path, help="CSV or PDF statement files")
    p.add_argument("--account", help="account name to tag these rows with, e.g. 'Chase Checking'")
    p.add_argument("--tracker", type=Path, default=DEFAULT_TRACKER,
                   help=f"tracking workbook (default: {DEFAULT_TRACKER.name} next to this script)")
    p.add_argument("--rules", type=Path, default=DEFAULT_RULES,
                   help="starter category rules, used only when creating a new tracker")
    p.add_argument("--init", action="store_true", help="create an empty tracker and exit")
    p.add_argument("--flip-sign", action="store_true",
                   help="negate amounts (for credit cards that show purchases as positive)")
    p.add_argument("--day-first", action="store_true", help="dates are DD/MM rather than MM/DD")
    p.add_argument("--date-format", help="explicit strptime format for CSV dates, e.g. %%d.%%m.%%Y")
    p.add_argument("--year", type=int, help="year for PDF dates printed without one")
    p.add_argument("--recategorize", action="store_true",
                   help="re-apply Categories rules to rows still marked Uncategorized")
    p.add_argument("--dry-run", action="store_true", help="show parsed rows without saving")
    args = p.parse_args(argv)

    if args.init:
        if args.tracker.exists():
            p.error(f"{args.tracker} already exists")
        new_workbook(load_rules_csv(args.rules)).save(args.tracker)
        print(f"Created {args.tracker}")
        return 0
    if not args.statements and not args.recategorize:
        p.error("give at least one statement file (or --init / --recategorize)")
    if args.statements and not args.account:
        p.error("--account is required when importing statements")
    for s in args.statements:
        if not s.exists():
            p.error(f"{s} not found")
    try:
        ingest(args.statements, args.account or "", args.tracker, args)
    except ValueError as e:
        print(f"error: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
